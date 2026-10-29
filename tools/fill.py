import sys, copy
sys.path.insert(0,'.')
from content import APP, S
import docx
d = docx.Document('/home/user/work/hmg/dl/form.docx')
def set_par(p, text):
    runs = p.runs
    if not runs: p.add_run(text); return
    runs[0].text = text
    for r in runs[1:]: r.text = ''
for p in d.paragraphs:
    t = p.text
    if t.startswith('참가 부문:'):
        set_par(p, '참가 부문: □ 대학(원)생    ■ 일반(스타트업 포함)')
    elif t.startswith('앱명 / 프로젝트명:'):
        set_par(p, f'앱명 / 프로젝트명: {APP}')
    elif '앱명 / 프로젝트명: ____' in t:
        set_par(p, t.split('앱명 / 프로젝트명:')[0] + f'앱명 / 프로젝트명: {APP}')
for i in range(1, 9):
    cell = d.tables[i].cell(0, 0)
    lines = S[i].split('\n')
    p0 = cell.paragraphs[0]
    set_par(p0, lines[0])
    for ln in lines[1:]:
        np = cell.add_paragraph(ln)
        if p0.style is not None: np.style = p0.style
        if p0.runs and np.runs:
            np.runs[0].font.size = p0.runs[0].font.size
            np.runs[0].font.name = p0.runs[0].font.name
d.save('/home/user/work/hmg/build/골든타임_접수양식_기획서_초안.docx')
print('saved')
